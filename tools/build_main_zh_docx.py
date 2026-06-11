#!/usr/bin/env python3
"""Build thesis-style DOCX from paper/main_zh.tex.

The LaTeX source remains the authority. This script only prepares a DOCX-friendly
copy for Pandoc, generates a static workflow diagram to replace the TikZ figure,
and then applies Chinese thesis-style Word formatting.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
SRC = PAPER / "main_zh.tex"
BUILD_DIR = PAPER / "build_docx"
WORKFLOW_IMG = PAPER / "figures" / "workflow_route.png"
PREPARED = BUILD_DIR / "main_zh_docx_source.tex"
OUT = PAPER / "main_zh.docx"

CHINESE_FONT = "宋体"
HEADING_FONT = "黑体"
LATIN_FONT = "Times New Roman"


def font_path() -> str | None:
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for p in candidates:
        if Path(p).exists():
            return p
    return None


def draw_centered(draw: ImageDraw.ImageDraw, box, text: str, font, fill=(20, 20, 20)):
    x0, y0, x1, y1 = box
    lines = text.split("\\n")
    line_heights = []
    max_w = 0
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        max_w = max(max_w, bbox[2] - bbox[0])
        line_heights.append(bbox[3] - bbox[1])
    total_h = sum(line_heights) + (len(lines) - 1) * 6
    y = y0 + ((y1 - y0) - total_h) / 2
    for line, h in zip(lines, line_heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        w = bbox[2] - bbox[0]
        draw.text((x0 + ((x1 - x0) - w) / 2, y), line, font=font, fill=fill)
        y += h + 6


def arrow(draw: ImageDraw.ImageDraw, start, end, fill=(55, 75, 95)):
    draw.line([start, end], fill=fill, width=5)
    sx, sy = start
    ex, ey = end
    import math
    angle = math.atan2(ey - sy, ex - sx)
    size = 18
    pts = [
        (ex, ey),
        (ex - size * math.cos(angle - math.pi / 6), ey - size * math.sin(angle - math.pi / 6)),
        (ex - size * math.cos(angle + math.pi / 6), ey - size * math.sin(angle + math.pi / 6)),
    ]
    draw.polygon(pts, fill=fill)


def make_workflow_image() -> None:
    WORKFLOW_IMG.parent.mkdir(parents=True, exist_ok=True)
    W, H = 1800, 1120
    im = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(im)
    fp = font_path()
    font = ImageFont.truetype(fp, 44) if fp else ImageFont.load_default()
    title_font = ImageFont.truetype(fp, 52) if fp else ImageFont.load_default()
    small_font = ImageFont.truetype(fp, 34) if fp else ImageFont.load_default()

    draw.text((W / 2, 55), "FG-MEE 板壳参数化仿真与验证路线", font=title_font, fill=(20, 35, 55), anchor="mm")

    boxes = {
        "mat": ((110, 160, 610, 310), "材料混合律\\nBaTiO3 / CoFe2O4", (221, 236, 255)),
        "fg": ((1190, 160, 1690, 310), "FG 分布 + 孔隙\\nU / V / X / O / P", (221, 236, 255)),
        "fem": ((110, 430, 610, 580), "MATLAB 板壳 FEM\\n八节点全耦合", (224, 247, 228)),
        "batch": ((1190, 430, 1690, 580), "静力扫描\\n2 mm 耦合验证", (224, 247, 228)),
        "comsol": ((110, 700, 610, 850), "COMSOL 独立验证\\n10 层 CSV 分域", (255, 238, 210)),
        "dyn": ((1190, 700, 1690, 850), "动力分析\\nNewmark + 模态降阶", (255, 238, 210)),
        "report": ((650, 900, 1150, 1050), "结果汇总\\n误差 < 5% / 孔隙复核", (255, 220, 220)),
    }
    for key, (box, label, color) in boxes.items():
        draw.rounded_rectangle(box, radius=28, fill=color, outline=(80, 100, 120), width=4)
        draw_centered(draw, box, label, font)

    def center(name):
        x0, y0, x1, y1 = boxes[name][0]
        return ((x0 + x1) / 2, (y0 + y1) / 2)

    arrow(draw, (610, 235), (1190, 235))
    arrow(draw, (360, 310), (360, 430))
    arrow(draw, (1440, 310), (1440, 430))
    arrow(draw, (610, 505), (1190, 505))
    arrow(draw, (360, 580), (360, 700))
    arrow(draw, (1440, 580), (1440, 700))
    arrow(draw, (610, 775), (650, 975))
    arrow(draw, (1190, 775), (1150, 975))

    draw.text((W/2, 1090), "注：COMSOL 验证用于独立校核；含孔隙代表行若超过 5% 则进入模型构型复核。", font=small_font, fill=(70, 70, 70), anchor="mm")
    im.save(WORKFLOW_IMG, quality=95)



def latex_math_to_plain(expr: str) -> str:
    """Best-effort inline LaTeX math to readable Word text."""
    out = expr.strip()
    # Common wrappers and text commands.
    for _ in range(6):
        out = re.sub(r"\\(?:text|mathrm|mathbf|mathit|bm|boldsymbol)\{([^{}]*)\}", r"\1", out)
        out = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", out)
    replacements = {
        r"\times": "×", r"\cdot": "·", r"\to": "→", r"\rightarrow": "→",
        r"\left": "", r"\right": "", r"\mathrm": "", r"\quad": " ", r"\;": " ", r"\,": "", r"\!": "",
        r"\%": r"\%", r"\&": "&", r"\_": "_",
        r"\alpha": "α", r"\beta": "β", r"\gamma": "γ", r"\kappa": "κ", r"\lambda": "λ",
        r"\theta": "θ", r"\omega": "ω", r"\xi": "ξ", r"\rho": "ρ", r"\chi": "χ",
        r"\Delta": "Δ", r"\Phi": "Φ", r"\Psi": "Ψ", r"\Omega": "Ω",
        r"\dot": "", r"\ddot": "", r"\bar": "",
        r"\sim": "~", r"\in": "∈", r"\le": "≤", r"\ge": "≥", r"\pm": "±", r"\approx": "≈",
        r"\textwidth": "textwidth",
    }
    for k, v in replacements.items():
        out = out.replace(k, v)
    out = re.sub(r"_\{([^{}]+)\}", r"_\1", out)
    out = re.sub(r"\^\{([^{}]+)\}", r"^\1", out)
    out = out.replace("{", "").replace("}", "")
    out = re.sub(r"\\[a-zA-Z]+", "", out)
    out = re.sub(r"\s+", " ", out).strip()
    return out


def inline_math_to_text(text: str) -> str:
    text = re.sub(r"(?<!\\)\$(?!\$)(.+?)(?<!\\)\$", lambda m: latex_math_to_plain(m.group(1)), text, flags=re.S)
    text = text.replace("BaTiO_3", "BaTiO₃").replace("CoFe_2O_4", "CoFe₂O₄")
    text = text.replace("BaTiO3", "BaTiO₃").replace("CoFe2O4", "CoFe₂O₄")
    return text


def simplify_docx_tables(text: str) -> str:
    """Make LaTeX table environments easier for Pandoc DOCX conversion."""
    text = re.sub(
        r"\\begin\{tabularx\}\{\\textwidth\}\{[^\n]*\}",
        r"\\begin{tabular}{p{1.3cm}p{4.1cm}p{4.1cm}p{4.6cm}}",
        text,
    )
    text = text.replace(r"\end{tabularx}", r"\end{tabular}")
    text = text.replace(r"\arraybackslash", "")
    return text

def preprocess_tex() -> None:
    text = SRC.read_text(encoding="utf-8")
    BUILD_DIR.mkdir(parents=True, exist_ok=True)

    text = text.replace("\\tableofcontents", "")
    text = text.replace("\\newpage", "")
    text = re.sub(r"\\textcolor\{[^{}]+\}\{([^{}]*)\}", r"\1", text)
    text = text.replace("\\SI", "")
    text = inline_math_to_text(text)
    text = text.replace(">6%", "大于 6\\%").replace("<5%", "小于 5\\%").replace(">6\\%", "大于 6\\%").replace("<5\\%", "小于 5\\%")
    text = simplify_docx_tables(text)

    # Replace TikZ workflow figure with a static PNG for DOCX.
    workflow_replacement = """\\begin{figure}[H]
\\centering
\\includegraphics[width=0.90\\textwidth]{figures/workflow_route.png}
\\caption{总体技术路线流程图}
\\label{fig:workflow}
\\end{figure}"""
    text = re.sub(
        r"\\begin\{figure\}\[H\]\s*\\centering\s*\\begin\{tikzpicture\}.*?\\caption\{总体技术路线流程图\}\s*\\label\{fig:workflow\}\s*\\end\{figure\}",
        lambda _m: workflow_replacement,
        text,
        flags=re.S,
    )

    # Pandoc can keep labels, but Chinese Word text is cleaner without raw labels.
    text = re.sub(r"\\label\{[^{}]+\}", "", text)
    text = re.sub(r"~?\\ref\{[^{}]+\}", "", text)
    text = re.sub(r"式~?\\eqref\{[^{}]+\}", "相关方程", text)
    text = re.sub(r"\\eqref\{[^{}]+\}", "相关方程", text)

    # Use simpler section labels for Word, no source comments.
    text = re.sub(r"^% ?={5,}.*$", "", text, flags=re.M)
    text = re.sub(r"^%.*$", "", text, flags=re.M)

    PREPARED.write_text(text, encoding="utf-8")


def run_pandoc() -> None:
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise RuntimeError("pandoc not found")
    resource_path = ":".join([
        str(PAPER),
        str(ROOT),
        str(ROOT / "reports" / "2026-05-28-porous-static" / "figures"),
        str(ROOT / "reports" / "2026-06-01-porous-dynamic" / "figures"),
    ])
    cmd = [
        pandoc,
        str(PREPARED),
        "--from", "latex+raw_tex",
        "--to", "docx",
        "--resource-path", resource_path,
        "--wrap=none",
        "--standalone",
        "-o", str(OUT),
    ]
    subprocess.run(cmd, cwd=ROOT, check=True)


def set_run_font(run, name=LATIN_FONT, east_asia=CHINESE_FONT, size=None, bold=None, color=None):
    run.font.name = name
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    else:
        run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def set_style_font(style, east_asia, latin, size, bold=None):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    rPr = style._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)


def set_cell_border(cell, **kwargs):
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, attrs in kwargs.items():
        tag = "w:{}".format(edge)
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def apply_three_line_table(table):
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            if ri == 0:
                set_cell_border(cell,
                    top={"val": "single", "sz": "12", "color": "000000"},
                    bottom={"val": "single", "sz": "6", "color": "000000"},
                    left={"val": "nil"}, right={"val": "nil"})
            elif ri == len(table.rows) - 1:
                set_cell_border(cell,
                    top={"val": "nil"}, bottom={"val": "single", "sz": "12", "color": "000000"},
                    left={"val": "nil"}, right={"val": "nil"})
            else:
                set_cell_border(cell,
                    top={"val": "nil"}, bottom={"val": "nil"},
                    left={"val": "nil"}, right={"val": "nil"})
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.first_line_indent = Pt(0)
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                para.paragraph_format.line_spacing = Pt(16)
                for run in para.runs:
                    set_run_font(run, size=10.5)


def paragraph_text(para):
    return "".join(run.text for run in para.runs).strip()


def postprocess_docx() -> None:
    doc = Document(OUT)

    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)

    styles = doc.styles
    set_style_font(styles["Normal"], CHINESE_FONT, LATIN_FONT, 12)
    pf = styles["Normal"].paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(22)
    pf.first_line_indent = Pt(24)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for name, east, size, before, after, align in [
        ("Title", HEADING_FONT, 16, 18, 18, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 1", HEADING_FONT, 16, 24, 24, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", HEADING_FONT, 14, 12, 12, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", HEADING_FONT, 12, 6, 6, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        if name in styles:
            set_style_font(styles[name], east, LATIN_FONT, size, bold=False)
            p = styles[name].paragraph_format
            p.first_line_indent = Pt(0)
            p.space_before = Pt(before)
            p.space_after = Pt(after)
            p.alignment = align

    if "Caption" in styles:
        set_style_font(styles["Caption"], CHINESE_FONT, LATIN_FONT, 10.5)
        styles["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styles["Caption"].paragraph_format.first_line_indent = Pt(0)

    table_caption_no = 0
    image_caption_no = 0
    for para in doc.paragraphs:
        original_style_name = para.style.name
        txt = paragraph_text(para)
        if not txt:
            continue
        # Pandoc metadata/title paragraphs may not receive Title style.
        if txt == "Abstract":
            para.text = "摘要"
            txt = "摘要"
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
        if original_style_name == "Table Caption" and not txt.startswith("表"):
            table_caption_no += 1
            para.text = f"表{table_caption_no} {txt}"
            txt = para.text.strip()
        elif original_style_name == "Image Caption" and not txt.startswith("图"):
            image_caption_no += 1
            para.text = f"图{image_caption_no} {txt}"
            txt = para.text.strip()
        if txt.startswith("含孔隙功能梯度磁"):
            para.style = styles["Title"]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
        elif para.style.name in {"Table Caption", "Image Caption"} or txt.startswith("图") or txt.startswith("表"):
            # Captions. Prefix numbering if Pandoc emitted raw caption text.
            try:
                para.style = styles["Caption"]
            except Exception:
                pass
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(6)
        elif para.style.name.startswith("Heading"):
            para.paragraph_format.first_line_indent = Pt(0)
        elif txt in {"作者", "2026", "摘要", "关键词"}:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if txt in {"作者", "2026"} else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.first_line_indent = Pt(0)
        for run in para.runs:
            size = None
            if para.style.name == "Title":
                size = 16
                set_run_font(run, east_asia=HEADING_FONT, size=size, bold=True)
            elif para.style.name.startswith("Heading 1"):
                set_run_font(run, east_asia=HEADING_FONT, size=16)
            elif para.style.name.startswith("Heading 2"):
                set_run_font(run, east_asia=HEADING_FONT, size=14)
            elif para.style.name.startswith("Heading 3"):
                set_run_font(run, east_asia=HEADING_FONT, size=12)
            elif para.style.name == "Caption" or txt.startswith(("图", "表")):
                set_run_font(run, size=10.5)
            else:
                set_run_font(run, size=12)

    for table in doc.tables:
        apply_three_line_table(table)
        # Wider, readable tables.
        table.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            table.autofit = True
        except Exception:
            pass

    # Header/footer with concise report marker and page number field.
    header = sec.header.paragraphs[0]
    header.text = "含孔隙功能梯度磁--电--弹性板仿真报告"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        set_run_font(run, size=10.5)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = ""
    run = footer.add_run("第 ")
    set_run_font(run, size=10.5)
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    fld_text = OxmlElement('w:t')
    fld_text.text = '1'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    r = footer.add_run()
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_sep)
    r._r.append(fld_text)
    r._r.append(fld_end)
    run2 = footer.add_run(" 页")
    set_run_font(run2, size=10.5)

    doc.save(OUT)


def verify_docx() -> None:
    doc = Document(OUT)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    table_texts.append(t)
    all_text = "\n".join(paras + table_texts)
    with zipfile.ZipFile(OUT) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    print(f"DOCX={OUT}")
    print(f"paragraphs={len(paras)}")
    print(f"tables={len(doc.tables)}")
    print(f"images={len(media)}")
    checks = [
        "功能梯度材料分布与含孔隙板壳分层模型示意图",
        "补充 MATLAB",
        "COMSOL 验证",
        "含孔隙 COMSOL 代表验证",
        "孔隙模式主导刚度退化",
    ]
    for c in checks:
        ok = c in all_text
        print(f"check:{c}={'OK' if ok else 'MISSING'}")
        if not ok:
            raise RuntimeError(f"missing key text: {c}")
    if len(doc.tables) < 10:
        raise RuntimeError("too few tables")
    if len(media) < 4:
        raise RuntimeError("too few embedded images")


def main() -> None:
    make_workflow_image()
    preprocess_tex()
    run_pandoc()
    postprocess_docx()
    verify_docx()


if __name__ == "__main__":
    main()
